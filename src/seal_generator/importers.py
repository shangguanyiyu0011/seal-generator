"""Import tabular company data from CSV and Excel files."""

import csv
from dataclasses import dataclass
from io import StringIO
from pathlib import Path

from openpyxl import load_workbook
import xlrd


@dataclass(frozen=True, slots=True)
class TableData:
    headers: tuple[str, ...]
    rows: tuple[tuple[object, ...], ...]


def _normalise(rows: list[list[object]], preview_limit: int | None, has_header: bool = True) -> TableData:
    if not rows or not any(value is not None and str(value).strip() for row in rows for value in row):
        raise ValueError("\u8868\u683c\u6ca1\u6709\u53ef\u7528\u7684\u8868\u5934")

    width = max(map(len, rows))
    headers = tuple(
        str(rows[0][index]).strip()
        if has_header and index < len(rows[0]) and rows[0][index] is not None and str(rows[0][index]).strip()
        else f"列 {index + 1}"
        for index in range(width)
    )
    start = 1 if has_header else 0
    body = rows[start:] if preview_limit is None else rows[start : start + preview_limit]
    normalised_rows = tuple(
        tuple(row[:width] + [None] * max(0, width - len(row))) for row in body
    )
    return TableData(headers=headers, rows=normalised_rows)


def _load_csv(path: Path) -> list[list[object]]:
    raw_data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            decoded = raw_data.decode(encoding)
        except UnicodeDecodeError:
            continue
        with StringIO(decoded, newline="") as handle:
            return [list(row) for row in csv.reader(handle)]
    raise ValueError("\u65e0\u6cd5\u8bc6\u522b CSV \u6587\u4ef6\u7f16\u7801")


def _load_xlsx(path: Path, sheet_name: str | None = None) -> list[list[object]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[sheet_name] if sheet_name is not None else book.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()


def _load_xls(path: Path, sheet_name: str | None = None) -> list[list[object]]:
    book = xlrd.open_workbook(path, on_demand=True)
    try:
        sheet = book.sheet_by_name(sheet_name) if sheet_name is not None else book.sheet_by_index(0)
        return [sheet.row_values(index) for index in range(sheet.nrows)]
    finally:
        book.release_resources()


def worksheet_names(path: Path) -> tuple[tuple[str, ...], str | None]:
    """Return worksheet choices without reading all of their cell data."""
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            names = tuple(sheet.title for sheet in book.worksheets)
            active = book.active.title if book.active is not None else None
            return names, active if active in names else (names[0] if names else None)
        finally:
            book.close()
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            names = tuple(book.sheet_names())
            return names, names[0] if names else None
        finally:
            book.release_resources()
    return (), None


def load_table(path: Path, preview_limit: int | None = None, *, has_header: bool = True, sheet_name: str | None = None) -> TableData:
    """Load the active/first worksheet or CSV into an immutable table."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _load_csv(path)
    elif suffix in {".xlsx", ".xlsm"}:
        rows = _load_xlsx(path, sheet_name)
    elif suffix == ".xls":
        rows = _load_xls(path, sheet_name)
    else:
        raise ValueError("\u4e0d\u652f\u6301\u7684\u6587\u4ef6\u683c\u5f0f\uff0c\u8bf7\u9009\u62e9 Excel \u6216 CSV \u6587\u4ef6")
    return _normalise(rows, preview_limit, has_header)


def extract_column(table: TableData, column_index: int) -> list[object]:
    """Return every value from the selected, existing column."""
    if isinstance(column_index, bool) or not isinstance(column_index, int):
        raise ValueError("\u5217\u7d22\u5f15\u5fc5\u987b\u662f\u6574\u6570")
    if not 0 <= column_index < len(table.headers):
        raise IndexError("\u6240\u9009\u516c\u53f8\u540d\u79f0\u5217\u4e0d\u5b58\u5728")
    return [row[column_index] for row in table.rows]
