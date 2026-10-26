from pathlib import Path
from threading import Event

from PIL import Image
from openpyxl import Workbook
from PySide6.QtCore import QThread, QTimer

from seal_generator.batch import BatchProcessor, failed_records, merge_summaries, pending_records
from seal_generator.importers import load_table, worksheet_names, extract_column
from seal_generator.models import BatchSummary, NameRecord, RecordStatus, ProgressSnapshot
from seal_generator.ui.main_window import MainWindow
import seal_generator.ui.main_window as ui_module


NAMES = ["北京华夏有限公司", "北京华夏印有限公司", "北京华夏印章有限公司"]


def ready(index, name):
    return NameRecord(index, name, name, len(name), RecordStatus.READY)


def test_headerless_csv_and_ragged_columns_never_discard_first_company(tmp_path):
    path = tmp_path / "names.csv"
    path.write_text(NAMES[0] + "\n" + NAMES[1] + ",备注\n", encoding="utf-8-sig")
    table = load_table(path, has_header=False)
    assert extract_column(table, 0) == NAMES[:2]
    assert table.rows[1][1] == "备注"
    assert len(table.headers) == 2


def test_explicit_worksheet_and_header_selection(tmp_path):
    path = tmp_path / "names.xlsx"
    book = Workbook()
    book.active.title = "说明"
    book.active.append(["使用说明"])
    sheet = book.create_sheet("名单")
    sheet.append(["公司名称"])
    for name in NAMES:
        sheet.append([name])
    book.save(path)
    assert worksheet_names(path) == (("说明", "名单"), "说明")
    assert extract_column(load_table(path, sheet_name="名单", has_header=True), 0) == NAMES
    assert extract_column(load_table(path, sheet_name="名单", has_header=False), 0) == ["公司名称", *NAMES]


def test_resume_and_retry_preserve_completed_paths_without_duplicate_files(tmp_path):
    processor = BatchProcessor(lambda n:Image.new("RGBA", (8, 8)))
    initial = processor.run([ready(i,n) for i,n in enumerate(NAMES)], tmp_path, lambda s:processor.cancel())
    assert [r.cleaned for r in pending_records(initial)] == NAMES[1:]
    resumed = BatchProcessor(lambda n:Image.new("RGBA", (8,8))).run(pending_records(initial), tmp_path)
    merged = merge_summaries(initial, resumed)
    assert not merged.cancelled
    assert merged.records[0].saved_path == initial.records[0].saved_path
    assert len(list(tmp_path.glob("*.png"))) == 3
    assert merged.final_snapshot == ProgressSnapshot(3,3,3,3,0,0)
    assert not pending_records(merged)
    merged.records[1].status = RecordStatus.FAILED
    retried = BatchProcessor(lambda n:Image.new("RGBA", (8,8))).run(failed_records(merged), tmp_path)
    final = merge_summaries(merged, retried)
    assert len(final.records) == 3
    assert final.records[0].saved_path == initial.records[0].saved_path
    assert final.final_snapshot.success == 3


def test_preview_is_async_coalesces_stale_selection_and_caches(qtbot, monkeypatch):
    window = MainWindow()
    qtbot.addWidget(window)
    entered, release = Event(), Event()
    calls, worker_threads = [], []
    def render(name):
        calls.append(name)
        worker_threads.append(QThread.currentThread() is window.thread())
        if name == NAMES[0]:
            entered.set()
            assert release.wait(5)
        return Image.new("RGBA", (16,16), (len(name), 0, 0, 255))
    monkeypatch.setattr(window, "_render", render)
    try:
        window.set_values(NAMES)
        qtbot.waitUntil(entered.is_set)
        ticks = []
        QTimer.singleShot(0, lambda:ticks.append(True))
        window.records_table.setCurrentCell(1,0)
        window.records_table.setCurrentCell(2,0)
        qtbot.waitUntil(lambda:bool(ticks))
        assert window.current_image is None
    finally:
        release.set()
    qtbot.waitUntil(lambda:window.current_image is not None)
    assert window.current_record.cleaned == NAMES[2]
    assert calls == [NAMES[0],NAMES[2]]
    assert worker_threads == [False, False]
    window.records_table.setCurrentCell(0,0)
    assert window.current_image is not None
    assert calls == [NAMES[0],NAMES[2]]


def test_new_source_resets_all_job_statistics(qtbot, monkeypatch):
    window = MainWindow()
    qtbot.addWidget(window)
    monkeypatch.setattr(window,"_render",lambda n:Image.new("RGBA",(8,8)))
    window._on_batch_finished(BatchSummary([NameRecord(0,NAMES[0],NAMES[0],8,RecordStatus.SUCCESS)]))
    window.set_values(NAMES[1:])
    assert window.batch_success_value.text() == "0"
    assert window.batch_failed_value.text() == "0"
    assert window.batch_total_value.text() == "0"
    assert not window.retry_button.isEnabled()
    assert not window.resume_button.isEnabled()


def test_gui_retry_merges_previous_success_rows(qtbot, monkeypatch, tmp_path):
    window = MainWindow()
    qtbot.addWidget(window)
    window.folder_edit.setText(str(tmp_path))
    success = NameRecord(0,NAMES[0],NAMES[0],8,RecordStatus.SUCCESS,actual_filename="old.png",saved_path=tmp_path/"old.png")
    failure = NameRecord(1,NAMES[1],NAMES[1],9,RecordStatus.FAILED,reason="保存失败")
    real_processor = BatchProcessor
    monkeypatch.setattr(ui_module,"BatchProcessor",lambda r:real_processor(lambda n:Image.new("RGBA",(8,8))))
    window._on_batch_finished(BatchSummary([success,failure]))
    window.retry_failed()
    qtbot.waitUntil(lambda:window._thread is None)
    assert window.results_table.rowCount() == 2
    assert window.results_table.item(0,2).text() == "old.png"
    assert window.batch_success_value.text() == "2"
    assert window.batch_total_value.text() == "2"


def test_import_uses_worker_and_limits_preview_without_limiting_records(qtbot, monkeypatch, tmp_path):
    window = MainWindow()
    qtbot.addWidget(window)
    monkeypatch.setattr(window,"_render",lambda n:Image.new("RGBA",(8,8)))
    path=tmp_path/"large.csv"
    path.write_text("\n".join(NAMES*50),encoding="utf-8-sig")
    window.load_import_file(path)
    assert window._importing
    assert not window.batch_button.isEnabled()
    qtbot.waitUntil(lambda:not window._importing)
    assert len(window.records) == 150
    assert window.records[0].cleaned == NAMES[0]
    assert window.import_preview.rowCount() == 100


def test_close_waits_for_preview_without_late_save(qtbot, monkeypatch, tmp_path):
    window=MainWindow()
    qtbot.addWidget(window)
    release=Event()
    def render(name):
        assert release.wait(5)
        return Image.new("RGBA",(8,8))
    monkeypatch.setattr(window,"_render",render)
    window.folder_edit.setText(str(tmp_path))
    try:
        window.set_values(NAMES[:1])
        window.save_current()
        window.close()
        assert window._close_pending
    finally:
        release.set()
    qtbot.waitUntil(lambda:window._preview_task is None)
    assert not list(tmp_path.glob("*.png"))


def test_header_and_sheet_controls_reload_real_workbook(qtbot, monkeypatch, tmp_path):
    path=tmp_path/"sheets.xlsx"
    book=Workbook()
    book.active.title="说明"
    book.active.append(["说明页"])
    sheet=book.create_sheet("公司")
    sheet.append(["公司名称"])
    for name in NAMES:sheet.append([name])
    book.save(path)
    window=MainWindow()
    qtbot.addWidget(window)
    monkeypatch.setattr(window,"_render",lambda n:Image.new("RGBA",(8,8)))
    window.load_import_file(path)
    qtbot.waitUntil(lambda:not window._importing)
    assert window.sheet_combo.count()==2
    window.sheet_combo.setCurrentText("公司")
    qtbot.waitUntil(lambda:not window._importing)
    assert len(window.records)==4
    window.header_check.setChecked(True)
    qtbot.waitUntil(lambda:not window._importing)
    assert [r.cleaned for r in window.records]==NAMES


def test_preview_finishing_during_folder_dialog_still_saves(qtbot, monkeypatch, tmp_path):
    window=MainWindow()
    qtbot.addWidget(window)
    monkeypatch.setattr(window,"_render",lambda n:Image.new("RGBA",(8,8)))
    window.set_values(NAMES[:1])
    def folder_dialog():
        qtbot.waitUntil(lambda:window.current_image is not None)
        return tmp_path
    monkeypatch.setattr(window,"_folder",folder_dialog)
    window.save_current()
    assert (tmp_path/f"{NAMES[0]}.png").is_file()


def test_chinese_spacing_follows_character_centres_not_stroke_widths():
    from seal_generator.layout import layout_name
    from seal_generator.resources import font_path
    for name in ["北京华夏有限公司", "某某市某某某中学", "一二三四五六七八"]:
        layout=layout_name(name,font_path())
        gaps=[a.angle_deg-b.angle_deg for a,b in zip(layout.glyphs,layout.glyphs[1:])]
        assert max(gaps)-min(gaps)<0.001
        # The new eight-character reference is about 30.5 degrees per centre,
        # and its endpoint centres descend below the horizontal diameter.
        assert 30<gaps[0]<31
        assert layout.glyphs[0].y>512
        assert layout.glyphs[-1].y>512
